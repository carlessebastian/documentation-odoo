#!/usr/bin/env python3
"""Importa facturas masivamente desde CSV (o XLSX) de forma idempotente.

Formato CSV requerido (cabeceras):
  ref,move_type,partner_vat,invoice_date,invoice_date_due,
  line_name,line_qty,line_price,line_tax_codes

Una fila por LINEA de factura. Las cabeceras de factura (ref, partner,
fechas) se duplican en cada fila; el script las agrupa por `ref` +
partner_vat + move_type.

Por defecto continua ante errores (devuelve `errors[]`). Con --atomic
crea todas en draft primero y solo postea si todas validan.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import OrderedDict
from typing import Any

from _common import OdooError
from odoo_client import OdooClient
from create_invoice import create_or_update_invoice


def _parse_csv(path: str) -> list[OrderedDict]:
    with open(path, encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        return [OrderedDict(row) for row in reader]


def _parse_xlsx(path: str) -> list[OrderedDict]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as e:
        raise OdooError(
            "openpyxl no instalado. `pip install openpyxl` o usar CSV."
        ) from e
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) for h in rows[0]]
    return [OrderedDict(zip(headers, [str(c) if c is not None else "" for c in row]))
            for row in rows[1:] if any(row)]


def _group_by_invoice(rows: list[OrderedDict]) -> list[dict]:
    grouped: OrderedDict[tuple, dict] = OrderedDict()
    for r in rows:
        key = (r["ref"], r["partner_vat"], r.get("move_type", "out_invoice"))
        if key not in grouped:
            grouped[key] = {
                "ref": r["ref"],
                "partner_vat": r["partner_vat"],
                "move_type": r.get("move_type", "out_invoice"),
                "invoice_date": r.get("invoice_date") or None,
                "invoice_date_due": r.get("invoice_date_due") or None,
                "lines": [],
            }
        tax_codes = [c.strip() for c in (r.get("line_tax_codes") or "").split(",") if c.strip()]
        grouped[key]["lines"].append({
            "name": r["line_name"],
            "qty": float(r["line_qty"]),
            "price_unit": float(r["line_price"]),
            "tax_codes": tax_codes,
        })
    return list(grouped.values())


def import_invoices(
    client: OdooClient,
    payloads: list[dict],
    *,
    post: bool = False,
    atomic: bool = False,
) -> dict[str, Any]:
    created: list[dict] = []
    updated: list[dict] = []
    skipped: list[dict] = []
    errors: list[dict] = []

    for p in payloads:
        try:
            p_with_post = dict(p, post=post and not atomic)
            result = create_or_update_invoice(client, p_with_post)
            status = result.get("status")
            if status == "created":
                created.append(result)
            elif status == "updated":
                updated.append(result)
            elif status == "already_posted":
                skipped.append(result)
            else:
                created.append(result)
        except OdooError as e:
            errors.append({"ref": p.get("ref"), "error": str(e)})
        except Exception as e:  # noqa: BLE001
            errors.append({"ref": p.get("ref"), "error": f"{type(e).__name__}: {e}"})

    if atomic and errors:
        # En modo atomic NO posteamos nada y reportamos los fallos.
        return {
            "atomic": True,
            "posted": [],
            "created_in_draft": created + updated,
            "errors": errors,
        }

    if atomic and post:
        ids = [r["id"] for r in created + updated]
        try:
            client.action("account.move", ids, "action_post")
            for r in created + updated:
                r["state"] = "posted"
        except OdooError as e:
            return {
                "atomic": True,
                "post_failed": True,
                "error": str(e),
                "created_in_draft": created + updated,
            }

    return {
        "total_input": len(payloads),
        "created": len(created),
        "updated": len(updated),
        "skipped_already_posted": len(skipped),
        "errors": errors,
        "details": {"created": created, "updated": updated, "skipped": skipped},
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", required=True)
    parser.add_argument("--format", choices=["csv", "xlsx"], default="csv")
    parser.add_argument("--post", action="store_true",
                        help="Postea tras crear (action_post)")
    parser.add_argument("--atomic", action="store_true",
                        help="Crea en draft, postea solo si todas pasan")
    parser.add_argument("--dry-run", action="store_true")
    ns = parser.parse_args()

    rows = _parse_csv(ns.file) if ns.format == "csv" else _parse_xlsx(ns.file)
    payloads = _group_by_invoice(rows)

    if ns.dry_run:
        json.dump(
            {"parsed_invoices": len(payloads), "preview": payloads[:3]},
            sys.stdout, indent=2, ensure_ascii=False,
        )
        print()
        return 0

    client = OdooClient()
    result = import_invoices(client, payloads, post=ns.post, atomic=ns.atomic)
    json.dump(result, sys.stdout, indent=2, ensure_ascii=False, default=str)
    print()
    return 0 if not result.get("errors") else 5


if __name__ == "__main__":
    try:
        sys.exit(main())
    except OdooError as e:
        print(f"OdooError: {e}", file=sys.stderr)
        sys.exit(2)
